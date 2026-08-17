# -*- coding: utf-8 -*-
"""Control de Entregas y Evidencias.

Visor documental por factura de cliente: qué se facturó, qué se entregó
realmente (movimientos de inventario validados menos devoluciones), qué
evidencias firmadas/selladas existen, y el seguimiento hasta el envío del
expediente a Administración.

No duplica información contable: todo lo fiscal se lee por campos
relacionados de account.move. Las cantidades se calculan desde los
movimientos reales de inventario ligados a las líneas de venta
(account.move.line.sale_line_ids → sale.order.line.move_ids). Los folios
de producción son los folios por línea de venta (delivery_folio) de este
mismo módulo: una orden de venta puede tener varios (S00300-1, S00300-2).
"""
import base64
import io
import logging
from datetime import timedelta

from odoo import api, fields, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

DELIVERY_STATES = [
    ('pending', 'Pendiente de entregar'),
    ('partial', 'Entrega parcial'),
    ('delivered', 'Entregado'),
    ('review', 'Requiere revisión'),
    ('cancelled', 'Cancelado'),
]

DOC_STATES = [
    ('no_evidence', 'Sin evidencia'),
    ('partial_evidence', 'Evidencia incompleta'),
    ('evidence_received', 'Evidencia recibida'),
    ('ready', 'Completo para Administración'),
    ('sent', 'Enviado a Administración'),
]


class DeliveryEvidenceControl(models.Model):
    _name = 'delivery.evidence.control'
    _description = 'Control de Entregas y Evidencias'
    _inherit = ['mail.thread', 'mail.activity.mixin']
    _order = 'invoice_date desc, id desc'
    _rec_name = 'name'

    # ------------------------------------------------------------------
    # Factura (todo relacionado, sin copiar datos contables)
    # ------------------------------------------------------------------
    move_id = fields.Many2one(
        'account.move', 'Factura', required=True, index=True,
        ondelete='restrict', domain=[('move_type', '=', 'out_invoice')],
    )
    name = fields.Char('Serie y Folio', related='move_id.name', store=True)
    company_id = fields.Many2one(related='move_id.company_id', store=True, index=True)
    partner_id = fields.Many2one(related='move_id.partner_id', store=True, string='Cliente')
    partner_code = fields.Char(
        'Código del cliente', related='move_id.partner_id.company_registry', store=True,
    )
    invoice_date = fields.Date(related='move_id.invoice_date', store=True, string='Fecha')
    currency_id = fields.Many2one(related='move_id.currency_id')
    amount_untaxed = fields.Monetary(related='move_id.amount_untaxed', string='Subtotal')
    amount_tax = fields.Monetary(related='move_id.amount_tax', string='Impuestos')
    amount_total = fields.Monetary(related='move_id.amount_total', string='Total')
    amount_total_signed = fields.Monetary(
        related='move_id.amount_total_signed', string='Total (moneda compañía)',
        currency_field='company_currency_id',
    )
    company_currency_id = fields.Many2one(related='move_id.company_currency_id')
    move_state = fields.Selection(related='move_id.state', string='Estado factura', store=True)

    # ------------------------------------------------------------------
    # Ventas, producción y logística (relaciones reales)
    # ------------------------------------------------------------------
    sale_order_ids = fields.Many2many(
        'sale.order', string='Órdenes de venta', compute='_compute_sale_links', store=True,
    )
    client_order_ref = fields.Char(
        'OC del cliente', compute='_compute_sale_links', store=True,
    )
    production_folios = fields.Char(
        'Folios de producción', compute='_compute_sale_links', store=True,
        help='Folios por línea de venta (multi-folio): cada consecutivo '
             'S00300-1, S00300-2… es un folio de producción independiente.',
    )
    picking_ids = fields.Many2many(
        'stock.picking', string='Entregas', compute='_compute_sale_links', store=True,
    )
    picking_count = fields.Integer(compute='_compute_counts')
    sale_count = fields.Integer(compute='_compute_counts')
    evidence_count = fields.Integer(compute='_compute_counts')

    # ------------------------------------------------------------------
    # Cantidades (calculadas de movimientos reales por _update_from_source)
    # ------------------------------------------------------------------
    line_ids = fields.One2many('delivery.evidence.control.line', 'control_id', 'Detalle')
    qty_invoiced = fields.Float('Cant. facturada', digits='Product Unit of Measure', readonly=True)
    qty_delivered = fields.Float('Cant. entregada (neta)', digits='Product Unit of Measure', readonly=True)
    qty_returned = fields.Float('Cant. devuelta', digits='Product Unit of Measure', readonly=True)
    qty_pending = fields.Float('Cant. pendiente', digits='Product Unit of Measure', readonly=True)
    delivered_pct = fields.Float('% entregado', readonly=True, aggregator='avg')
    delivery_state = fields.Selection(
        DELIVERY_STATES, 'Estado de entrega', default='pending',
        readonly=True, index=True, tracking=True,
    )
    review_reason = fields.Text('Motivo de revisión', readonly=True)

    # ------------------------------------------------------------------
    # Evidencias y seguimiento administrativo
    # ------------------------------------------------------------------
    evidence_ids = fields.One2many('delivery.evidence.document', 'control_id', 'Evidencias')
    doc_state = fields.Selection(
        DOC_STATES, 'Estado documental', default='no_evidence',
        readonly=True, index=True, tracking=True,
    )
    evidence_received_date = fields.Date('Fecha de evidencia', readonly=True)
    evidence_user_id = fields.Many2one('res.users', 'Recibió evidencia', readonly=True)
    responsible_id = fields.Many2one('res.users', 'Responsable', tracking=True)
    ready_exception = fields.Boolean('Completado con excepción', readonly=True)
    sent_date = fields.Datetime('Enviado a Administración', readonly=True)
    sent_user_id = fields.Many2one('res.users', 'Envió', readonly=True)
    notes = fields.Text('Observaciones')
    days_without_evidence = fields.Integer(
        'Días sin evidencia', compute='_compute_days_without_evidence',
        help='Días desde la fecha de factura sin evidencia validada.',
    )
    active = fields.Boolean(default=True)

    _sql_constraints = [
        ('move_uniq', 'unique(move_id)',
         'Ya existe un control de entregas y evidencias para esa factura.'),
    ]

    # ==================================================================
    # Cómputos ligeros
    # ==================================================================
    @api.depends('move_id.invoice_line_ids.sale_line_ids')
    def _compute_sale_links(self):
        for control in self:
            sale_lines = control.move_id.invoice_line_ids.sale_line_ids
            orders = sale_lines.order_id
            control.sale_order_ids = [(6, 0, orders.ids)]
            refs = [r for r in orders.mapped('client_order_ref') if r]
            if not refs and control.move_id.ref:
                refs = [control.move_id.ref]
            control.client_order_ref = ', '.join(dict.fromkeys(refs)) or False
            folios = [f for f in sale_lines.mapped('delivery_folio') if f]
            control.production_folios = ', '.join(dict.fromkeys(folios)) or False
            control.picking_ids = [(6, 0, orders.picking_ids.ids)]

    def _compute_counts(self):
        for control in self:
            control.picking_count = len(control.picking_ids)
            control.sale_count = len(control.sale_order_ids)
            control.evidence_count = len(control.evidence_ids)

    @api.depends('doc_state', 'invoice_date')
    def _compute_days_without_evidence(self):
        today = fields.Date.context_today(self)
        for control in self:
            if control.doc_state in ('no_evidence', 'partial_evidence') and control.invoice_date:
                control.days_without_evidence = (today - control.invoice_date).days
            else:
                control.days_without_evidence = 0

    # ==================================================================
    # Cálculo de cantidades desde los movimientos reales
    # ==================================================================
    def _sale_line_delivery_qty(self, sale_line, target_uom):
        """(entregado, devuelto) de una línea de venta, en la UdM destino.

        Solo movimientos validados (state='done'). Entrega = destino en
        ubicación de cliente; devolución = origen en ubicación de cliente.
        """
        delivered = returned = 0.0
        for move in sale_line.move_ids.filtered(lambda m: m.state == 'done'):
            qty = move.product_uom._compute_quantity(
                move.quantity, target_uom, rounding_method='HALF-UP')
            if move.location_dest_id.usage == 'customer':
                delivered += qty
            elif move.location_id.usage == 'customer':
                returned += qty
        return delivered, returned

    def _update_from_source(self):
        """Recalcula relaciones, cantidades y estado de entrega del control.

        No modifica facturas, entregas ni inventario: solo lee.
        """
        Line = self.env['delivery.evidence.control.line']
        for control in self:
            move = control.move_id
            control.line_ids.unlink()

            if move.state == 'cancel':
                control.write({
                    'delivery_state': 'cancelled',
                    'qty_invoiced': 0.0, 'qty_delivered': 0.0,
                    'qty_returned': 0.0, 'qty_pending': 0.0,
                    'delivered_pct': 0.0, 'review_reason': False,
                })
                continue

            # Notas de crédito publicadas ligadas a esta factura: reducen lo
            # facturado por producto para no inflar los totales.
            refunds = self.env['account.move'].search([
                ('move_type', '=', 'out_refund'),
                ('state', '=', 'posted'),
                ('reversed_entry_id', '=', move.id),
            ])
            refund_pool = {}
            for rline in refunds.invoice_line_ids:
                if rline.display_type == 'product' and rline.product_id:
                    key = rline.product_id.id
                    refund_pool[key] = refund_pool.get(key, 0.0) + rline.quantity

            reasons = []
            totals = {'inv': 0.0, 'del': 0.0, 'ret': 0.0, 'pen': 0.0}
            product_lines = move.invoice_line_ids.filtered(
                lambda l: l.display_type == 'product' and l.product_id)

            for inv_line in product_lines:
                uom = inv_line.product_uom_id or inv_line.product_id.uom_id
                qty_inv = inv_line.quantity
                # Aplica la nota de crédito disponible para este producto.
                pool = refund_pool.get(inv_line.product_id.id, 0.0)
                if pool > 0:
                    applied = min(pool, qty_inv)
                    qty_inv -= applied
                    refund_pool[inv_line.product_id.id] = pool - applied

                sale_lines = inv_line.sale_line_ids
                delivered = returned = 0.0
                needs_review = False
                line_reason = False

                if not sale_lines:
                    needs_review = True
                    line_reason = _(
                        'La línea de factura no está ligada a ninguna línea de '
                        'venta; la entrega no puede determinarse sin ambigüedad.')
                else:
                    for sl in sale_lines:
                        d, r = control._sale_line_delivery_qty(sl, uom)
                        delivered += d
                        returned += r
                        # Si la misma línea de venta está facturada en más de
                        # una factura publicada, la atribución de lo entregado
                        # a ESTA factura es ambigua: no se inventa un reparto.
                        other_invoices = sl.invoice_lines.move_id.filtered(
                            lambda m: m.move_type == 'out_invoice'
                            and m.state == 'posted' and m != move)
                        if other_invoices:
                            needs_review = True
                            line_reason = _(
                                'La línea de venta %(folio)s también está '
                                'facturada en %(others)s: lo entregado no puede '
                                'atribuirse sin ambigüedad a una sola factura.'
                            ) % {
                                'folio': sl.delivery_folio or sl.order_id.name,
                                'others': ', '.join(other_invoices.mapped('name')),
                            }

                net = delivered - returned
                pending = max(qty_inv - net, 0.0)
                Line.create({
                    'control_id': control.id,
                    'invoice_line_id': inv_line.id,
                    'sale_line_id': sale_lines[:1].id,
                    'product_id': inv_line.product_id.id,
                    'uom_id': uom.id,
                    'qty_invoiced': qty_inv,
                    'qty_delivered_raw': delivered,
                    'qty_returned': returned,
                    'qty_delivered_net': net,
                    'qty_pending': pending,
                    'needs_review': needs_review,
                    'review_reason': line_reason,
                })
                if needs_review and line_reason:
                    reasons.append(line_reason)
                totals['inv'] += qty_inv
                totals['del'] += net
                totals['ret'] += returned
                totals['pen'] += pending

            rounding = 0.001
            if reasons:
                state = 'review'
            elif totals['inv'] <= rounding and not product_lines:
                state = 'review'
                reasons.append(_('La factura no tiene líneas de producto.'))
            elif totals['pen'] <= rounding:
                state = 'delivered'
            elif totals['del'] > rounding:
                state = 'partial'
            else:
                state = 'pending'

            control.write({
                'qty_invoiced': totals['inv'],
                'qty_delivered': totals['del'],
                'qty_returned': totals['ret'],
                'qty_pending': totals['pen'],
                'delivered_pct': (100.0 * totals['del'] / totals['inv']) if totals['inv'] else 0.0,
                'delivery_state': state,
                'review_reason': '\n'.join(dict.fromkeys(reasons)) or False,
            })

    def _update_evidence_stage(self):
        """Etapas automáticas de evidencia (no toca ready/sent, que son acciones)."""
        for control in self:
            if control.doc_state in ('ready', 'sent'):
                continue
            validated = control.evidence_ids.filtered(lambda e: e.state == 'validated')
            if validated:
                control.doc_state = 'evidence_received'
                if not control.evidence_received_date:
                    control.evidence_received_date = fields.Date.context_today(control)
                    control.evidence_user_id = self.env.user
            elif control.evidence_ids:
                control.doc_state = 'partial_evidence'
            else:
                control.doc_state = 'no_evidence'

    # ==================================================================
    # Sincronización (usada por el wizard y por la publicación de facturas)
    # ==================================================================
    @api.model
    def _sync_from_moves(self, moves):
        """Crea/actualiza controles para las facturas dadas. Idempotente."""
        stats = {'created': 0, 'updated': 0, 'skipped': 0, 'review': 0}
        moves = moves.filtered(lambda m: m.move_type == 'out_invoice')
        if not moves:
            return stats
        existing = {
            c.move_id.id: c
            for c in self.with_context(active_test=False).search(
                [('move_id', 'in', moves.ids)])
        }
        for move in moves:
            if move.state != 'posted' and move.id not in existing:
                stats['skipped'] += 1
                continue
            control = existing.get(move.id)
            if control:
                control._update_from_source()
                stats['updated'] += 1
            else:
                control = self.create({'move_id': move.id})
                control._update_from_source()
                stats['created'] += 1
            if control.delivery_state == 'review':
                stats['review'] += 1
        return stats

    # ==================================================================
    # Acciones
    # ==================================================================
    def _is_manager(self):
        return self.env.user.has_group('restricciones_entregas.group_delivery_evidence_manager')

    def action_refresh(self):
        self._update_from_source()
        self._update_evidence_stage()
        return True

    def action_validate_evidence(self):
        self.ensure_one()
        if not self._is_manager():
            raise UserError(_('Solo el responsable de Entregas y Evidencias puede validar evidencias.'))
        drafts = self.evidence_ids.filtered(lambda e: e.state == 'draft')
        if not drafts:
            raise UserError(_('No hay evidencias pendientes de validar.'))
        drafts.action_validate()
        return True

    def action_mark_ready(self):
        self.ensure_one()
        if self.doc_state == 'sent':
            raise UserError(_('El control ya fue enviado a Administración.'))
        if not self.evidence_ids.filtered(lambda e: e.state == 'validated'):
            raise UserError(_('No se puede completar el expediente sin al menos una evidencia validada.'))
        exception = False
        if self.qty_pending > 0.001 or self.delivery_state != 'delivered':
            if not self._is_manager():
                raise UserError(_(
                    'La factura aún tiene cantidad pendiente de entregar. Solo el '
                    'responsable puede completar el expediente con una excepción justificada.'))
            if not (self.notes or '').strip():
                raise UserError(_(
                    'Para usar la excepción captura primero la justificación en Observaciones.'))
            exception = True
        self.write({'doc_state': 'ready', 'ready_exception': exception})
        body = _('Expediente marcado como completo para Administración.')
        if exception:
            body = _(
                'EXCEPCIÓN: expediente completado con %(qty)s pendiente de entregar. '
                'Justificación: %(notes)s'
            ) % {'qty': self.qty_pending, 'notes': self.notes}
        self.message_post(body=body)
        return True

    def action_mark_sent(self):
        if not self._is_manager():
            raise UserError(_('Solo el responsable puede marcar el envío a Administración.'))
        for control in self:
            if control.doc_state != 'ready':
                raise UserError(_(
                    'El control %s no está listo para Administración.') % control.name)
            control.write({
                'doc_state': 'sent',
                'sent_date': fields.Datetime.now(),
                'sent_user_id': self.env.user.id,
            })
            control.message_post(body=_('Enviado a Administración.'))
        return True

    def action_reopen(self):
        self.ensure_one()
        if not self._is_manager():
            raise UserError(_('Solo el responsable puede reabrir un control.'))
        self.write({'ready_exception': False, 'sent_date': False, 'sent_user_id': False})
        self._update_evidence_stage()
        self.message_post(body=_('Control reabierto por el responsable.'))
        return True

    def action_generate_report(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Generar relación'),
            'res_model': 'delivery.evidence.report.wizard',
            'view_mode': 'form',
            'target': 'new',
            'context': {'default_control_ids': [(6, 0, self.ids)]},
        }

    def action_open_invoice(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'account.move',
            'res_id': self.move_id.id,
            'view_mode': 'form',
        }

    def action_open_pickings(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Entregas'),
            'res_model': 'stock.picking',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.picking_ids.ids)],
        }

    def action_open_sales(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Órdenes de venta'),
            'res_model': 'sale.order',
            'view_mode': 'list,form',
            'domain': [('id', 'in', self.sale_order_ids.ids)],
        }

    def unlink(self):
        if any(c.doc_state == 'sent' for c in self) and not self._is_manager():
            raise UserError(_('Un control enviado a Administración no se puede eliminar.'))
        return super().unlink()

    # ==================================================================
    # API para la aplicación OWL (Centro de operación)
    # ==================================================================
    TAB_DOMAINS = {
        'all': [],
        'pending': [('delivery_state', 'in', ['pending', 'partial'])],
        'no_evidence': [('delivery_state', '=', 'delivered'),
                        ('doc_state', 'in', ['no_evidence', 'partial_evidence'])],
        'ready': [('doc_state', '=', 'ready')],
        'sent': [('doc_state', '=', 'sent')],
        'review': [('delivery_state', '=', 'review')],
    }

    @api.model
    def js_bootstrap(self):
        counts = {
            tab: self.search_count(domain)
            for tab, domain in self.TAB_DOMAINS.items()
        }
        return {
            'user_name': self.env.user.name,
            'is_manager': self.env.user.has_group(
                'restricciones_entregas.group_delivery_evidence_manager'),
            'counts': counts,
            'evidence_types': [
                {'value': v, 'label': l}
                for v, l in self.env['delivery.evidence.document']._fields['evidence_type'].selection
            ],
        }

    @api.model
    def js_list(self, tab='all', search='', limit=120):
        domain = list(self.TAB_DOMAINS.get(tab, []))
        if search:
            term = search.strip()
            domain += ['|', '|', '|', '|',
                       ('name', 'ilike', term),
                       ('partner_id', 'ilike', term),
                       ('client_order_ref', 'ilike', term),
                       ('production_folios', 'ilike', term),
                       ('sale_order_ids.name', 'ilike', term)]
        controls = self.search(domain, limit=limit, order='invoice_date desc, id desc')
        return [c._js_row() for c in controls]

    def _js_row(self):
        self.ensure_one()
        return {
            'id': self.id,
            'name': self.name or '',
            'partner': self.partner_id.name or '',
            'partner_code': self.partner_code or '',
            'date': self.invoice_date and self.invoice_date.strftime('%d/%m/%Y') or '',
            'amount_total': self.amount_total,
            'currency': self.currency_id.name or 'MXN',
            'qty_invoiced': self.qty_invoiced,
            'qty_delivered': self.qty_delivered,
            'qty_pending': self.qty_pending,
            'pct': round(self.delivered_pct, 1),
            'delivery_state': self.delivery_state,
            'doc_state': self.doc_state,
            'days': self.days_without_evidence,
            'folios': self.production_folios or '',
            'oc': self.client_order_ref or '',
            'sales': ', '.join(self.sale_order_ids.mapped('name')),
            'evidence_count': len(self.evidence_ids),
            'sent_date': self.sent_date and fields.Datetime.context_timestamp(
                self, self.sent_date).strftime('%d/%m/%Y') or '',
            'exception': self.ready_exception,
        }

    def js_detail(self):
        self.ensure_one()
        data = self._js_row()
        data.update({
            'amount_untaxed': self.amount_untaxed,
            'amount_tax': self.amount_tax,
            'move_state': self.move_state,
            'review_reason': self.review_reason or '',
            'notes': self.notes or '',
            'responsible': self.responsible_id.name or '',
            'evidence_received_date': self.evidence_received_date and
                self.evidence_received_date.strftime('%d/%m/%Y') or '',
            'lines': [{
                'id': l.id,
                'product': l.product_id.display_name or '',
                'folio': l.production_folio or '',
                'uom': l.uom_id.name or '',
                'qty_invoiced': l.qty_invoiced,
                'qty_delivered': l.qty_delivered_net,
                'qty_returned': l.qty_returned,
                'qty_pending': l.qty_pending,
                'review': l.needs_review,
                'review_reason': l.review_reason or '',
            } for l in self.line_ids],
            'evidences': [{
                'id': e.id,
                'type': e.evidence_type,
                'type_label': dict(e._fields['evidence_type'].selection)[e.evidence_type],
                'name': e.name,
                'file_name': e.file_name or '',
                'url': '/web/content?model=delivery.evidence.document&field=file'
                       f'&id={e.id}&download=true&filename={e.file_name or e.name}',
                'doc_date': e.doc_date and e.doc_date.strftime('%d/%m/%Y') or '',
                'uploaded_by': e.create_uid.name,
                'uploaded_at': fields.Datetime.context_timestamp(
                    e, e.create_date).strftime('%d/%m/%Y %H:%M'),
                'state': e.state,
                'validated_by': e.validated_by_id.name or '',
                'notes': e.notes or '',
            } for e in self.evidence_ids],
        })
        return data

    def js_add_evidence(self, vals):
        self.ensure_one()
        self.env['delivery.evidence.document'].create({
            'control_id': self.id,
            'evidence_type': vals.get('evidence_type') or 'remision_firmada',
            'name': vals.get('name') or vals.get('file_name') or _('Evidencia'),
            'file': vals['file'],
            'file_name': vals.get('file_name'),
            'doc_date': vals.get('doc_date') or False,
            'notes': vals.get('notes') or False,
        })
        return self.js_detail()

    def js_action(self, action):
        """Ejecuta una acción del flujo y regresa el detalle actualizado."""
        self.ensure_one()
        actions = {
            'refresh': self.action_refresh,
            'validate': self.action_validate_evidence,
            'ready': self.action_mark_ready,
            'sent': self.action_mark_sent,
            'reopen': self.action_reopen,
        }
        if action not in actions:
            raise UserError(_('Acción no reconocida.'))
        actions[action]()
        return self.js_detail()

    def js_set_notes(self, notes):
        self.ensure_one()
        self.notes = notes or False
        return True

    def js_validate_document(self, document_id):
        self.ensure_one()
        doc = self.evidence_ids.filtered(lambda d: d.id == document_id)
        doc.action_validate()
        return self.js_detail()

    @api.model
    def js_bulk_action(self, control_ids, action):
        """Aplica una acción a varios controles; reporta éxito/fallo por folio."""
        results = {'ok': [], 'failed': []}
        for control in self.browse(control_ids):
            try:
                with self.env.cr.savepoint():
                    control.js_action(action)
                results['ok'].append(control.name)
            except UserError as error:
                results['failed'].append({'name': control.name, 'reason': str(error)})
        return results

    @api.model
    def js_sync_recent(self, days=60):
        if not self.env.user.has_group('restricciones_entregas.group_delivery_evidence_manager'):
            raise UserError(_('Solo el responsable puede sincronizar facturas.'))
        date_from = fields.Date.context_today(self) - timedelta(days=days)
        moves = self.env['account.move'].search([
            ('move_type', '=', 'out_invoice'),
            ('state', '=', 'posted'),
            ('invoice_date', '>=', date_from),
        ])
        stats = self._sync_from_moves(moves)
        stats['total'] = len(moves)
        return stats

    @api.model
    def js_match_excel(self, file_b64, filename):
        """Empata un Excel contra los controles por folio de factura, orden de
        venta o folio de producción: sirve para palomear en lote lo que venga
        listado en cualquier layout de Excel."""
        content = base64.b64decode(file_b64)
        fname = (filename or '').lower()
        tokens = set()
        if fname.endswith('.xlsx'):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for value in row:
                        if isinstance(value, str) and 2 < len(value.strip()) <= 64:
                            tokens.add(value.strip().lower())
        elif fname.endswith('.xls'):
            try:
                import xlrd
            except ImportError:
                raise UserError(_('El servidor no puede leer .xls; guarda el archivo como .xlsx.'))
            wb = xlrd.open_workbook(file_contents=content)
            for sheet in wb.sheets():
                for r in range(sheet.nrows):
                    for c in range(sheet.ncols):
                        value = sheet.cell_value(r, c)
                        if isinstance(value, str) and 2 < len(value.strip()) <= 64:
                            tokens.add(value.strip().lower())
        else:
            raise UserError(_('Sube un archivo .xlsx o .xls.'))

        index = {}
        for control in self.search([]):
            identifiers = [control.name or '']
            identifiers += (control.production_folios or '').split(', ')
            identifiers += control.sale_order_ids.mapped('name')
            for identifier in identifiers:
                if identifier:
                    index.setdefault(identifier.strip().lower(), control)

        matched = self.browse()
        for token in tokens:
            control = index.get(token)
            if control:
                matched |= control
        return {
            'matched': [c._js_row() for c in matched.sorted(
                key=lambda c: (c.invoice_date or fields.Date.today(), c.id))],
            'cells_scanned': len(tokens),
        }


class DeliveryEvidenceControlLine(models.Model):
    _name = 'delivery.evidence.control.line'
    _description = 'Detalle por producto del control de entregas'
    _order = 'id'

    control_id = fields.Many2one(
        'delivery.evidence.control', required=True, index=True, ondelete='cascade')
    company_id = fields.Many2one(related='control_id.company_id', store=True)
    invoice_line_id = fields.Many2one('account.move.line', 'Línea de factura', readonly=True)
    sale_line_id = fields.Many2one('sale.order.line', 'Línea de venta', readonly=True)
    production_folio = fields.Char(
        'Folio de producción', related='sale_line_id.delivery_folio', store=True)
    product_id = fields.Many2one('product.product', 'Producto', readonly=True)
    uom_id = fields.Many2one('uom.uom', 'UdM', readonly=True)
    qty_invoiced = fields.Float('Facturado', digits='Product Unit of Measure', readonly=True)
    qty_delivered_raw = fields.Float('Entregado', digits='Product Unit of Measure', readonly=True)
    qty_returned = fields.Float('Devuelto', digits='Product Unit of Measure', readonly=True)
    qty_delivered_net = fields.Float('Entregado neto', digits='Product Unit of Measure', readonly=True)
    qty_pending = fields.Float('Pendiente', digits='Product Unit of Measure', readonly=True)
    needs_review = fields.Boolean('Requiere revisión', readonly=True)
    review_reason = fields.Text('Motivo', readonly=True)


class DeliveryEvidenceDocument(models.Model):
    _name = 'delivery.evidence.document'
    _description = 'Evidencia de entrega'
    _order = 'create_date desc'

    control_id = fields.Many2one(
        'delivery.evidence.control', required=True, index=True, ondelete='cascade')
    company_id = fields.Many2one(related='control_id.company_id', store=True)
    evidence_type = fields.Selection([
        ('remision_firmada', 'Remisión firmada'),
        ('remision_sellada', 'Remisión sellada'),
        ('acuse', 'Acuse de entrega'),
        ('foto', 'Fotografía'),
        ('documento', 'Documento adicional'),
        ('otro', 'Otro'),
    ], 'Tipo', required=True, default='remision_firmada')
    name = fields.Char('Nombre', required=True)
    file = fields.Binary('Archivo', attachment=True, required=True)
    file_name = fields.Char('Nombre del archivo')
    doc_date = fields.Date('Fecha del documento')
    notes = fields.Char('Observaciones')
    state = fields.Selection([
        ('draft', 'Cargada'),
        ('validated', 'Validada'),
    ], 'Validación', default='draft', required=True)
    validated_by_id = fields.Many2one('res.users', 'Validó', readonly=True)
    validated_date = fields.Datetime('Fecha de validación', readonly=True)

    @api.model_create_multi
    def create(self, vals_list):
        docs = super().create(vals_list)
        for doc in docs:
            doc.control_id.message_post(body=_(
                'Evidencia cargada: %(type)s — %(name)s'
            ) % {'type': dict(doc._fields['evidence_type'].selection)[doc.evidence_type],
                 'name': doc.name})
        docs.control_id._update_evidence_stage()
        return docs

    def action_validate(self):
        manager = self.env.user.has_group(
            'restricciones_entregas.group_delivery_evidence_manager')
        if not manager:
            raise UserError(_('Solo el responsable puede validar evidencias.'))
        for doc in self.filtered(lambda d: d.state == 'draft'):
            doc.write({
                'state': 'validated',
                'validated_by_id': self.env.user.id,
                'validated_date': fields.Datetime.now(),
            })
            doc.control_id.message_post(body=_('Evidencia validada: %s') % doc.name)
        self.control_id._update_evidence_stage()
        return True

    def unlink(self):
        manager = self.env.user.has_group(
            'restricciones_entregas.group_delivery_evidence_manager')
        controls = self.control_id
        for doc in self:
            if doc.state == 'validated' and not manager:
                raise UserError(_('Una evidencia validada solo la puede eliminar el responsable.'))
            doc.control_id.message_post(body=_('Evidencia eliminada: %s') % doc.name)
        res = super().unlink()
        controls._update_evidence_stage()
        return res
